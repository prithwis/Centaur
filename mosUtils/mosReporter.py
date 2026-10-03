# ==================================================
#  C E N T A U R  ::  MOS
#  Middle East Oil Security Simulator
#
#  Prithwis Mukerjee | 2026
#  Action -> Consequence -> Reaction -> ...
# ==================================================

from mos import *


class PressReporter_Agent:

    # ========================================================
    # Initialise PressReporter
    # ========================================================

    def __init__(
        self,
        client,
        model,
        agent_id,
        generic_role
    ):
        self.client = client
        self.model = model
        self.agent_id = agent_id
        self.generic_role = generic_role


    # ========================================================
    # Convert one WorldState into text
    # ========================================================

    def _worldStateText(self, world):

        text = []

        text.append(f"WORLDSTATE: {world.state_id}")
        text.append("-" * 60)

        for variable in world.variables:

            text.append(f"\nVariable: {variable.name}")
            text.append(f"Index: {variable.index}")
            text.append(f"Description: {variable.description}")
            text.append(f"Direction: {variable.direction}")
            text.append(f"Status: {variable.status}")
            text.append(f"Rationale: {variable.rationale}")

            if variable.conditions:
                text.append("Conditions:")

                for condition in variable.conditions:
                    text.append(f"- {condition}")

        return "\n".join(text)


    # ========================================================
    # Convert one Event into text
    # ========================================================

    def _eventText(self, event):

        text = []

        text.append(f"EVENT: {event.event_id}")
        text.append("-" * 60)

        text.append(f"Actor: {event.actor}")
        text.append(f"Action: {event.action}")
        text.append(f"Rationale: {event.rationale}")
        text.append(
            f"Selection Weight: {event.selection_weight}"
        )

        return "\n".join(text)


    # ========================================================
    # Convert complete Scenario into text
    #
    # W0 -> E0 -> W1 -> E1 -> ... -> En -> W(n+1)
    # ========================================================

    def _scenarioText(self, W, E):

        # A completed Scenario must have one more WorldState
        # than Events.
        if len(W) != len(E) + 1:
            raise ValueError(
                "Incomplete Scenario: "
                "len(W) must equal len(E) + 1."
            )

        text = []

        text.append("MOS SIMULATION SCENARIO")
        text.append("=" * 60)

        text.append(f"Events: {len(E)}")
        text.append(f"WorldStates: {len(W)}")

        # ----------------------------------------------------
        # Initial WorldState
        # ----------------------------------------------------

        text.append("\n\nINITIAL WORLDSTATE")
        text.append("=" * 60)
        text.append(self._worldStateText(W[0]))

        # ----------------------------------------------------
        # Event / resulting WorldState sequence
        # ----------------------------------------------------

        for i, event in enumerate(E):

            text.append("\n\n")
            text.append(self._eventText(event))

            text.append("\n\nRESULTING WORLDSTATE")
            text.append("=" * 60)
            text.append(self._worldStateText(W[i + 1]))

        return "\n".join(text)


    # ========================================================
    # Build Reporting Role
    #
    # Combines the generic PressReporter role with the
    # report-specific Narrative or Structured prompt.
    # ========================================================

    def _buildRole(self, prompt):

        return f"""
{self.generic_role}

{prompt}

Use only the supplied MOS simulation Scenario.

Do not introduce external facts, events or assumptions.
Do not alter Event IDs, WorldState IDs, numerical indices
or selection weights.
"""


    # ========================================================
    # Narrative Report
    #
    # Produces a flowing human-readable account of the Scenario,
    # explaining Events, consequences and major turning points.
    # ========================================================

    def narrative(self, W, E, prompt):

        scenario = self._scenarioText(W, E)
        role = self._buildRole(prompt)

        print(
            f"PressReporter working ---------- "
            f"{self.agent_id} | Narrative"
        )

        return callLLM(
            self.client,
            role,
            scenario,
            self.model
        )


    # ========================================================
    # Structured Report
    #
    # Produces a structured Markdown representation suitable
    # for subsequent machine-based analysis.
    # ========================================================

    def structured(self, W, E, prompt):

        scenario = self._scenarioText(W, E)
        role = self._buildRole(prompt)

        print(
            f"PressReporter working ---------- "
            f"{self.agent_id} | Structured"
        )

        return callLLM(
            self.client,
            role,
            scenario,
            self.model
        )


# ============================================================

# Microsoft DOCX generator 


# ============================================================
# DOCX HELPER FUNCTIONS
# ============================================================

def _addFormattedText(paragraph, text):
    """Convert simple Markdown bold/italic into Word formatting."""

    import re

    pattern = r'(\*\*.*?\*\*|\*.*?\*)'
    parts = re.split(pattern, text)

    for part in parts:

        if not part:
            continue

        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True

        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True

        else:
            paragraph.add_run(part)


def _addPageNumber(paragraph):
    """Insert an automatic Word PAGE field."""

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


# ============================================================
# Helper: Add WorldState Evolution Table
# ============================================================

def _addWorldStateTable(doc, W, E):
    """
    Add a compact table showing the evolution of WorldVariable
    indices across the Scenario.

    Rows    : WorldStates W0, W1, W2, ...
    Columns : WorldVariables
    Event   : Event that produced each WorldState
              (W0 is the Initial State)
    """

    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

    if not W:
        return

    # --------------------------------------------------------
    # Table heading
    # --------------------------------------------------------

    p = doc.add_paragraph()

    run = p.add_run("WorldState Evolution")
    run.bold = True
    run.font.size = Pt(13)

    p = doc.add_paragraph()

    run = p.add_run(
        "All values are simulation indices normalised to 100 "
        "at the initial WorldState (W0)."
    )

    run.italic = True
    run.font.size = Pt(9)


    # --------------------------------------------------------
    # Column headings
    # --------------------------------------------------------

    # Short headings keep the table readable in portrait mode.
    variableHeaders = [
        "Oil",
        "Hormuz",
        "India",
        "Sanctions",
        "Allies",
        "Iran",
        "US-Iran",
        "Conflict"
    ]

    headers = [
        "State",
        "Event"
    ] + variableHeaders

    table = doc.add_table(
        rows=1,
        cols=len(headers)
    )

    table.style = "Table Grid"


    # --------------------------------------------------------
    # Header row
    # --------------------------------------------------------

    headerCells = table.rows[0].cells

    for i, heading in enumerate(headers):

        headerCells[i].text = heading
        headerCells[i].vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

        p = headerCells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for run in p.runs:
            run.bold = True
            run.font.size = Pt(8)


    # --------------------------------------------------------
    # WorldState rows
    # --------------------------------------------------------

    for stateNo, world in enumerate(W):

        cells = table.add_row().cells

        # WorldState
        cells[0].text = world.state_id

        # Event producing this WorldState
        if stateNo == 0:
            cells[1].text = "Initial"
        else:
            cells[1].text = E[stateNo - 1].event_id

        # Variable indices
        for j, variable in enumerate(world.variables):

            cells[j + 2].text = str(variable.index)


        # Format row
        for cell in cells:

            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for run in p.runs:
                run.font.size = Pt(8)


    # --------------------------------------------------------
    # Event Key
    # --------------------------------------------------------

    doc.add_paragraph()

    p = doc.add_paragraph()

    run = p.add_run("Event Key")
    run.bold = True
    run.font.size = Pt(10)

    for event in E:

        p = doc.add_paragraph()

        p.paragraph_format.space_after = Pt(2)

        run = p.add_run(
            f"{event.event_id} — {event.actor}: "
        )

        run.bold = True
        run.font.size = Pt(9)

        run = p.add_run(event.action)

        run.font.size = Pt(9)
    


# ============================================================
# NARRATIVE DOCX CREATOR
# ============================================================

def createNarrativeDocx(
    narrativeText,
    W,
    E,
    scenarioID="Scenario",
    title="MOS Scenario Report",
    outputDir="/content"
):

    import os
    import re

    from datetime import datetime
    from zoneinfo import ZoneInfo

    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn


    # --------------------------------------------------------
    # Filename
    # --------------------------------------------------------

    now = datetime.now(ZoneInfo("Asia/Kolkata"))

    filename = (
        f"MOS_{scenarioID}_{now.strftime('%y%m%d_%H%M%S')}.docx"
    )

    filepath = os.path.join(outputDir, filename)


    # --------------------------------------------------------
    # Create document
    # --------------------------------------------------------

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)


    # --------------------------------------------------------
    # Normal text
    # --------------------------------------------------------

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11.5)

    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.10


    # --------------------------------------------------------
    # Heading styles
    # --------------------------------------------------------

    headingStyles = [
    ("Title", 21),
    ("Heading 1", 17),
    ("Heading 2", 14),
    ("Heading 3", 12)
    ]

    for styleName, fontSize in headingStyles:

        style = doc.styles[styleName]

        style.font.name = "Aptos Display"
        style.font.size = Pt(fontSize)
        style.font.bold = True


    # --------------------------------------------------------
    # Header
    # --------------------------------------------------------

    header = section.header

    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = p.add_run(
        "MOS   |   MIDDLE EAST OIL SECURITY SIMULATION"
    )

    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.bold = True


    # --------------------------------------------------------
    # Footer
    # --------------------------------------------------------

    footer = section.footer

    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    generated = now.strftime(
        "%d %b %Y  %H:%M:%S"
    )

    run = p.add_run(
        f"Generated on {generated}   |   "
        f"Prithwis Mukerjee   |   "
        f"Page "
    )

    run.font.name = "Aptos"
    run.font.size = Pt(8)

    _addPageNumber(p)


    # --------------------------------------------------------
    # Document title
    # --------------------------------------------------------

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = p.add_run(title)

    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)

    p = doc.add_paragraph()

    run = p.add_run(
        f"Scenario: {scenarioID}"
    )

    run.font.name = "Aptos"
    run.font.size = Pt(10)

    doc.add_paragraph()

    # --------------------------------------------------------
    # Simulation Note
    # --------------------------------------------------------

    trigger = E[0]
    actors = list(dict.fromkeys(event.actor for event in E))
    actorText = " and ".join(actors)

    # E0 is the initial trigger; subsequent Events are
    # simulation-generated responses.
    steps = len(E) - 1

    p = doc.add_paragraph()

    # Indent the Simulation Note
    p.paragraph_format.left_indent = Inches(0.30)
    p.paragraph_format.right_indent = Inches(0.20)

    # Add vertical bar on the left
    pPr = p._p.get_or_add_pPr()

    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")

    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), "808080")

    pBdr.append(left)
    pPr.append(pBdr)

    run = p.add_run(
        f"This is an LLM-based geopolitical simulation, not an account "
        f"of actual events. The simulation begins with the hypothetical "
        f"trigger: {trigger.action} It then evolves through {steps} "
        f"simulation-generated steps involving {actorText}, with each "
        f"action changing the simulated world and influencing the next response. "
        f"Numerical values shown in the report are simulation indices, not "
        f"real-world measurements: each variable is normalised to 100 at the "
        f"start of the simulation, and subsequent values indicate movement "
        f"relative to that initial baseline."
    )

    run.italic = True
    run.font.name = "Aptos"
    run.font.size = Pt(10.5)

    doc.add_paragraph()


    # --------------------------------------------------------
    # Convert Narrative Markdown to Word
    # --------------------------------------------------------

    for rawLine in narrativeText.splitlines():

        line = rawLine.strip()

        if not line:
            continue

        if line.startswith("# "):

            doc.add_heading(
                line[2:].strip(),
                level=1
            )

        elif line.startswith("## "):

            doc.add_heading(
                line[3:].strip(),
                level=2
            )

        elif line.startswith("### "):

            doc.add_heading(
                line[4:].strip(),
                level=3
            )

        elif line.startswith(">"):

            p = doc.add_paragraph()

            p.paragraph_format.left_indent = Inches(0.3)

            run = p.add_run(
                line[1:].strip()
            )

            run.italic = True

        elif re.match(r"^[-*]\s+", line):

            text = re.sub(
                r"^[-*]\s+",
                "",
                line
            )

            p = doc.add_paragraph(
                style="List Bullet"
            )

            _addFormattedText(p, text)

        elif re.match(r"^\d+\.\s+", line):

            text = re.sub(
                r"^\d+\.\s+",
                "",
                line
            )

            p = doc.add_paragraph(
                style="List Number"
            )

            _addFormattedText(p, text)

        elif re.match(r"^-{3,}$", line):

            continue

        else:

            p = doc.add_paragraph()

            _addFormattedText(
                p,
                line
            )


    # --------------------------------------------------------
    # WorldState Evolution
    # --------------------------------------------------------

    _addWorldStateTable(doc, W, E)


    # --------------------------------------------------------
    # Project Attribution
    # --------------------------------------------------------

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    run = p.add_run(
        "Generated by Prithwis Mukerjee — "
        "linkedin.com/in/prithwis/ — "
        "under the Centaur Project."
    )

    run.italic = True
    run.font.size = Pt(9)


    # --------------------------------------------------------
    # Centaur Logo
    # --------------------------------------------------------

    logoPath = "/content/CentaurLogo.png"

    if os.path.exists(logoPath):

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        run = p.add_run()
        run.add_picture(
            logoPath,
            width=Inches(0.8)
        )


    # --------------------------------------------------------
    # Document metadata
    # --------------------------------------------------------

    doc.core_properties.title = title

    doc.core_properties.subject = (
        "MOS Middle East Oil Security Simulation"
    )

    doc.core_properties.author = (
        "Prithwis Mukerjee"
    )

    doc.core_properties.keywords = (
        f"MOS, Centaur, ZeitWorld, Chanakya, "
        f"Scenario Simulation, {scenarioID}"
    )


    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    doc.save(filepath)

    print(
        f"Word report created: {filepath}"
    )

    return filepath