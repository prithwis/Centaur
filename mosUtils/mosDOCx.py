# ==================================================
#  C E N T A U R  ::  MOS
#  Middle East Oil Security Simulator
#
#  Prithwis Mukerjee | 2026
#  Action -> Consequence -> Reaction -> ...
# ==================================================

from mos import *

# Microsoft DOCX generator 


# ============================================================
# DOCX HELPER FUNCTIONS
# ============================================================

def _addFormattedText(paragraph, text):
    """Convert simple Markdown bold/italic into Word formatting."""

    import re

    pattern = r'(\*\*.*?\*\*|\*.*?\*)'
    parts = re.split(pattern, text)

    for part in parts:

        if not part:
            continue

        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True

        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True

        else:
            paragraph.add_run(part)


def _addPageNumber(paragraph):
    """Insert an automatic Word PAGE field."""

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


# ============================================================
# Helper: Add WorldState Evolution Table
# ============================================================

def _addWorldStateTable(doc, W, E):
    """
    Add a compact table showing the evolution of WorldVariable
    indices across the Scenario.

    Rows    : WorldStates W0, W1, W2, ...
    Columns : WorldVariables
    Event   : Event that produced each WorldState
              (W0 is the Initial State)
    """

    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

    if not W:
        return

    # --------------------------------------------------------
    # Table heading
    # --------------------------------------------------------

    p = doc.add_paragraph()

    run = p.add_run("WorldState Evolution")
    run.bold = True
    run.font.size = Pt(13)

    p = doc.add_paragraph()

    run = p.add_run(
        "All values are simulation indices normalised to 100 "
        "at the initial WorldState (W0)."
    )

    run.italic = True
    run.font.size = Pt(9)


    # --------------------------------------------------------
    # Column headings
    # --------------------------------------------------------

    variableHeaders = {
        "Oil Price": "Oil",
        "Hormuz Oil Flow": "Hormuz",
        "Regional Diplomatic Engagement": "Diplomacy",
        "Economic Pressure on Iran": "Econ Pressure",
        "US/Regional Partner Cohesion": "US Partners",
        "Iran Domestic Stability": "Iran Stability",
        "US-Iran Tension": "US-Iran",
        "Regional Armed Conflict": "Conflict"
    }

    headers = ["State", "Event"] + [
        variableHeaders.get(variable.name, variable.name)
        for variable in W[0].variables
]

    table = doc.add_table(
        rows=1,
        cols=len(headers)
    )

    table.style = "Table Grid"


    # --------------------------------------------------------
    # Header row
    # --------------------------------------------------------

    headerCells = table.rows[0].cells

    for i, heading in enumerate(headers):

        headerCells[i].text = heading
        headerCells[i].vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

        p = headerCells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for run in p.runs:
            run.bold = True
            run.font.size = Pt(8)


    # --------------------------------------------------------
    # WorldState rows
    # --------------------------------------------------------

    for stateNo, world in enumerate(W):

        cells = table.add_row().cells

        # WorldState
        cells[0].text = world.state_id

        # Event producing this WorldState
        if stateNo == 0:
            cells[1].text = "Initial"
        else:
            cells[1].text = E[stateNo - 1].event_id

        # Variable indices
        for j, variable in enumerate(world.variables):

            cells[j + 2].text = str(variable.index)


        # Format row
        for cell in cells:

            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for run in p.runs:
                run.font.size = Pt(8)


    # --------------------------------------------------------
    # Event Key
    # --------------------------------------------------------

    doc.add_paragraph()

    p = doc.add_paragraph()

    run = p.add_run("Event Key")
    run.bold = True
    run.font.size = Pt(10)

    for event in E:

        p = doc.add_paragraph()

        p.paragraph_format.space_after = Pt(2)

        run = p.add_run(
            f"{event.event_id} — {event.actor}: "
        )

        run.bold = True
        run.font.size = Pt(9)

        run = p.add_run(event.action)

        run.font.size = Pt(9)
    


# ============================================================
# NARRATIVE DOCX CREATOR
# ============================================================

def createNarrativeDocx(
    narrativeText,
    W,
    E,
    scenarioID="Scenario",
    title="MOS Scenario Report",
    outputDir="/content"
):

    import os
    import re

    from datetime import datetime
    from zoneinfo import ZoneInfo

    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn


    # --------------------------------------------------------
    # Filename
    # --------------------------------------------------------

    now = datetime.now(ZoneInfo("Asia/Kolkata"))

    filename = (
        f"MOS_{scenarioID}_{now.strftime('%y%m%d_%H%M%S')}.docx"
    )

    filepath = os.path.join(outputDir, filename)


    # --------------------------------------------------------
    # Create document
    # --------------------------------------------------------

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)


    # --------------------------------------------------------
    # Normal text
    # --------------------------------------------------------

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11.5)

    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.10


    # --------------------------------------------------------
    # Heading styles
    # --------------------------------------------------------

    headingStyles = [
    ("Title", 21),
    ("Heading 1", 17),
    ("Heading 2", 14),
    ("Heading 3", 12)
    ]

    for styleName, fontSize in headingStyles:

        style = doc.styles[styleName]

        style.font.name = "Aptos Display"
        style.font.size = Pt(fontSize)
        style.font.bold = True


    # --------------------------------------------------------
    # Header
    # --------------------------------------------------------

    header = section.header

    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = p.add_run(
        "MOS   |   MIDDLE EAST OIL SECURITY SIMULATION"
    )

    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.bold = True


    # --------------------------------------------------------
    # Footer
    # --------------------------------------------------------

    footer = section.footer

    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    generated = now.strftime(
        "%d %b %Y  %H:%M:%S"
    )

    run = p.add_run(
        f"Generated on {generated}   |   "
        f"Prithwis Mukerjee   |   "
        f"Page "
    )

    run.font.name = "Aptos"
    run.font.size = Pt(8)

    _addPageNumber(p)


    # --------------------------------------------------------
    # Document title
    # --------------------------------------------------------

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = p.add_run(title)

    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)

    p = doc.add_paragraph()

    run = p.add_run(
        f"Scenario: {scenarioID}"
    )

    run.font.name = "Aptos"
    run.font.size = Pt(10)

    doc.add_paragraph()

    # --------------------------------------------------------
    # Simulation Note
    # --------------------------------------------------------

    trigger = E[0]
    actors = list(dict.fromkeys(event.actor for event in E))
    actorText = " and ".join(actors)

    # E0 is the initial trigger; subsequent Events are
    # simulation-generated responses.
    steps = len(E) - 1

    p = doc.add_paragraph()

    # Indent the Simulation Note
    p.paragraph_format.left_indent = Inches(0.30)
    p.paragraph_format.right_indent = Inches(0.20)

    # Add vertical bar on the left
    pPr = p._p.get_or_add_pPr()

    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")

    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), "808080")

    pBdr.append(left)
    pPr.append(pBdr)

    run = p.add_run(
        f"This is an LLM-based geopolitical simulation, not an account "
        f"of actual events. The simulation begins with the hypothetical "
        f"trigger: {trigger.action} It then evolves through {steps} "
        f"simulation-generated steps involving {actorText}, with each "
        f"action changing the simulated world and influencing the next response. "
        f"Numerical values shown in the report are simulation indices, not "
        f"real-world measurements: each variable is normalised to 100 at the "
        f"start of the simulation, and subsequent values indicate movement "
        f"relative to that initial baseline."
    )

    run.italic = True
    run.font.name = "Aptos"
    run.font.size = Pt(10.5)

    doc.add_paragraph()


    # --------------------------------------------------------
    # Convert Narrative Markdown to Word
    # --------------------------------------------------------

    for rawLine in narrativeText.splitlines():

        line = rawLine.strip()

        if not line:
            continue

        if line.startswith("# "):

            doc.add_heading(
                line[2:].strip(),
                level=1
            )

        elif line.startswith("## "):

            doc.add_heading(
                line[3:].strip(),
                level=2
            )

        elif line.startswith("### "):

            doc.add_heading(
                line[4:].strip(),
                level=3
            )

        elif line.startswith(">"):

            p = doc.add_paragraph()

            p.paragraph_format.left_indent = Inches(0.3)

            run = p.add_run(
                line[1:].strip()
            )

            run.italic = True

        elif re.match(r"^[-*]\s+", line):

            text = re.sub(
                r"^[-*]\s+",
                "",
                line
            )

            p = doc.add_paragraph(
                style="List Bullet"
            )

            _addFormattedText(p, text)

        elif re.match(r"^\d+\.\s+", line):

            text = re.sub(
                r"^\d+\.\s+",
                "",
                line
            )

            p = doc.add_paragraph(
                style="List Number"
            )

            _addFormattedText(p, text)

        elif re.match(r"^-{3,}$", line):

            continue

        else:

            p = doc.add_paragraph()

            _addFormattedText(
                p,
                line
            )


    # --------------------------------------------------------
    # WorldState Evolution
    # --------------------------------------------------------

    _addWorldStateTable(doc, W, E)


    # --------------------------------------------------------
    # Project Attribution
    # --------------------------------------------------------

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    run = p.add_run(
        "Generated by Prithwis Mukerjee — "
        "linkedin.com/in/prithwis/ — "
        "under the Centaur Project."
    )

    run.italic = True
    run.font.size = Pt(9)


    # --------------------------------------------------------
    # Centaur Logo
    # --------------------------------------------------------

    logoPath = "/content/CentaurLogo.png"

    if os.path.exists(logoPath):

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        run = p.add_run()
        run.add_picture(
            logoPath,
            width=Inches(0.8)
        )


    # --------------------------------------------------------
    # Document metadata
    # --------------------------------------------------------

    doc.core_properties.title = title

    doc.core_properties.subject = (
        "MOS Middle East Oil Security Simulation"
    )

    doc.core_properties.author = (
        "Prithwis Mukerjee"
    )

    doc.core_properties.keywords = (
        f"MOS, Centaur, ZeitWorld, Chanakya, "
        f"Scenario Simulation, {scenarioID}"
    )


    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    doc.save(filepath)

    print(
        f"Word report created: {filepath}"
    )

    return filepath